"""Extract text content from various file types for RAG ingestion.

Supports PDF, Word documents, Excel spreadsheets, and images (placeholder).
"""

import logging
from pathlib import Path

logger = logging.getLogger("renovation_assistant")


def extract_text_from_pdf(file_path: Path) -> str:
    """Extract text content from a PDF file.

    Args:
        file_path: Path to the PDF file.

    Returns:
        Extracted text content.

    Raises:
        ImportError: If pypdf is not installed.
        Exception: If PDF extraction fails.
    """
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        logger.exception("pypdf not installed")
        raise ImportError("pypdf library is required for PDF extraction") from exc

    try:
        reader = PdfReader(file_path)
        text_parts: list[str] = []

        for page_num, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text()
            if page_text.strip():
                text_parts.append(f"[Page {page_num}]\n{page_text.strip()}")

        if not text_parts:
            logger.warning("No text extracted from PDF: %s", file_path)
            return ""

        return "\n\n".join(text_parts)
    except Exception as exc:
        logger.exception("Failed to extract text from PDF: %s", file_path)
        raise RuntimeError(f"PDF extraction failed: {exc}") from exc


def extract_text_from_docx(file_path: Path) -> str:
    """Extract text content from a Word document.

    Args:
        file_path: Path to the DOCX file.

    Returns:
        Extracted text content.

    Raises:
        ImportError: If python-docx is not installed.
        Exception: If DOCX extraction fails.
    """
    try:
        from docx import Document
    except ImportError as exc:
        logger.exception("python-docx not installed")
        raise ImportError("python-docx library is required for Word document extraction") from exc

    try:
        doc = Document(file_path)
        paragraphs: list[str] = []

        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        if not paragraphs:
            logger.warning("No text extracted from DOCX: %s", file_path)
            return ""

        return "\n\n".join(paragraphs)
    except Exception as exc:
        logger.exception("Failed to extract text from DOCX: %s", file_path)
        raise RuntimeError(f"DOCX extraction failed: {exc}") from exc


def extract_text_from_excel(file_path: Path) -> str:
    """Extract text content from an Excel spreadsheet.

    Args:
        file_path: Path to the XLSX file.

    Returns:
        Extracted text content in table format.

    Raises:
        ImportError: If openpyxl is not installed.
        Exception: If Excel extraction fails.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        logger.exception("openpyxl not installed")
        raise ImportError("openpyxl library is required for Excel extraction") from exc

    try:
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            sheet_contents: list[str] = []

            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                rows: list[str] = []

                for row in sheet.iter_rows(values_only=True):
                    # Filter out completely empty rows
                    row_values = [str(cell) if cell is not None else "" for cell in row]
                    if any(val.strip() for val in row_values):
                        rows.append(" | ".join(row_values))

                if rows:
                    sheet_text = f"[Sheet: {sheet_name}]\n" + "\n".join(rows)
                    sheet_contents.append(sheet_text)

            if not sheet_contents:
                logger.warning("No data extracted from Excel: %s", file_path)
                return ""

            return "\n\n".join(sheet_contents)
        finally:
            workbook.close()
    except Exception as exc:
        logger.exception("Failed to extract text from Excel: %s", file_path)
        raise RuntimeError(f"Excel extraction failed: {exc}") from exc


def extract_text_from_csv(file_path: Path) -> str:
    """Extract text content from a CSV file.

    Args:
        file_path: Path to the CSV file.

    Returns:
        Extracted text content in table format.

    Raises:
        Exception: If CSV extraction fails.
    """
    import csv

    try:
        rows: list[str] = []

        with file_path.open("r", encoding="utf-8", errors="ignore") as csv_file:
            reader = csv.reader(csv_file)
            for row in reader:
                if any(cell.strip() for cell in row):
                    rows.append(" | ".join(row))

        if not rows:
            logger.warning("No data extracted from CSV: %s", file_path)
            return ""

        return "\n".join(rows)
    except Exception as exc:
        logger.exception("Failed to extract text from CSV: %s", file_path)
        raise RuntimeError(f"CSV extraction failed: {exc}") from exc


def extract_text_from_image(file_path: Path) -> str:
    """Extract text or description from an image file.

    Note: This is a placeholder. For production, integrate OCR (tesseract/pytesseract)
    or LLM vision capabilities. For now, returns a description placeholder.

    Args:
        file_path: Path to the image file.

    Returns:
        Placeholder description text.
    """
    # TODO: Integrate OCR or vision model for actual text extraction
    logger.info("Image file %s processed with placeholder (OCR not implemented)", file_path)
    return (
        f"[Image: {file_path.name}]\n"
        "Image content processing requires OCR or vision model integration."
    )


def extract_file_content(file_path: Path, file_type: str) -> str:
    """Extract text content from a file based on its MIME type.

    Args:
        file_path: Path to the file.
        file_type: MIME type of the file.

    Returns:
        Extracted text content.

    Raises:
        ValueError: If file type is not supported.
        Exception: If extraction fails.
    """
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    # Map MIME type to extraction function
    if file_type == "application/pdf":
        return extract_text_from_pdf(file_path)
    elif file_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return extract_text_from_docx(file_path)
    elif file_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return extract_text_from_excel(file_path)
    elif file_type == "text/csv":
        return extract_text_from_csv(file_path)
    elif file_type.startswith("image/"):
        return extract_text_from_image(file_path)
    else:
        raise ValueError(f"Unsupported file type for content extraction: {file_type}")
